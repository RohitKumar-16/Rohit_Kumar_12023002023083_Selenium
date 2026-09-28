import os
import json
import time
from datetime import datetime

from selenium import webdriver
from selenium.webdriver.common.by import By
from selenium.webdriver.chrome.service import Service
from selenium.webdriver.chrome.options import Options
from selenium.webdriver.support.ui import WebDriverWait
from selenium.webdriver.support import expected_conditions as EC
from selenium.common.exceptions import (
    TimeoutException,
    NoAlertPresentException,
    ElementClickInterceptedException,
)


# Paths / constants
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
SCREENSHOT_DIR = os.path.join(BASE_DIR, "screenshots")
TEST_DATA_JSON = os.path.join(BASE_DIR, "test_data.json")
TEST_DATA_XLSX = os.path.join(BASE_DIR, "test_data.xlsx")  # optional
WAIT_TIMEOUT = 15


CHROMEDRIVER_PATH = None

os.makedirs(SCREENSHOT_DIR, exist_ok=True)



# Read test data from JSON / Excel
def load_test_data():
    """Uses test_data.xlsx if it exists, otherwise test_data.json."""
    if os.path.exists(TEST_DATA_XLSX):
        return _load_from_excel(TEST_DATA_XLSX)
    with open(TEST_DATA_JSON, "r", encoding="utf-8") as f:
        return json.load(f)


def _load_from_excel(path):
    from openpyxl import load_workbook

    ws = load_workbook(path).active
    headers = [c.value for c in ws[1]]
    row = [c.value for c in ws[2]]
    record = dict(zip(headers, row))
    return {
        "base_url": record["base_url"],
        "login": {
            "email": record["email"],
            "password": record["password"],
            "name": record["name"],
        },
        "search": {
            "product_search_term": record["product_search_term"],
            "product_name_to_add": record["product_name_to_add"],
            "quantity": int(record["quantity"]),
        },
    }


# Launch browse
def launch_browser(headless=False):
    options = Options()
    if headless:
        options.add_argument("--headless=new")
    options.add_argument("--start-maximized")
    options.add_argument("--disable-notifications")
    options.add_argument("--disable-popup-blocking")

    if CHROMEDRIVER_PATH:
        driver = webdriver.Chrome(service=Service(CHROMEDRIVER_PATH), options=options)
    else:
        driver = webdriver.Chrome(options=options)

    driver.implicitly_wait(2)
    return driver


# Screenshots
def take_screenshot(driver, name):
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    filename = f"{name.replace(' ', '_').lower()}_{timestamp}.png"
    filepath = os.path.join(SCREENSHOT_DIR, filename)
    driver.save_screenshot(filepath)
    return filepath


# Popups / alerts and safe clicking
def dismiss_browser_alert_if_present(driver, accept=True):
    try:
        WebDriverWait(driver, 3).until(EC.alert_is_present())
        alert = driver.switch_to.alert
        text = alert.text
        alert.accept() if accept else alert.dismiss()
        print(f"Handled a browser alert: '{text}'")
        return text
    except (TimeoutException, NoAlertPresentException):
        return None


def safe_click(driver, element):
    driver.execute_script("arguments[0].scrollIntoView({block: 'center'});", element)
    time.sleep(0.3)
    try:
        element.click()
    except ElementClickInterceptedException:
        driver.execute_script("arguments[0].click();", element)


def close_add_to_cart_modal(driver):
    try:
        continue_btn = WebDriverWait(driver, WAIT_TIMEOUT).until(
            EC.element_to_be_clickable((By.XPATH, "//button[text()='Continue Shopping']"))
        )
        safe_click(driver, continue_btn)
        print("Dismissed 'Added to cart' popup.")
    except TimeoutException:
        print("No 'Added to cart' popup appeared.")


# Login (falls back to signup if the account does not exist yet)
def is_logged_in(driver):
    return len(driver.find_elements(By.XPATH, "//a[contains(., 'Logged in as')]")) > 0


def login_or_signup(driver, data):
    login = data["login"]
    wait = WebDriverWait(driver, WAIT_TIMEOUT)
    driver.get(f"{data['base_url']}/login")

    email_field = wait.until(EC.presence_of_element_located((By.NAME, "email")))
    password_field = driver.find_element(By.NAME, "password")
    login_btn = driver.find_element(By.CSS_SELECTOR, "button[data-qa='login-button']")

    email_field.clear()
    email_field.send_keys(login["email"])
    password_field.clear()
    password_field.send_keys(login["password"])
    safe_click(driver, login_btn)

    time.sleep(1.5)  # let the page redirect or show an error

    error_visible = driver.find_elements(
        By.XPATH, "//p[contains(text(), 'incorrect') or contains(text(), 'Incorrect')]"
    )
    if error_visible:
        print("No account found for this email - signing up instead.")
        _signup_new_account(driver, login, wait)


def _signup_new_account(driver, login, wait):
    driver.find_element(By.NAME, "name").clear()
    driver.find_element(By.NAME, "name").send_keys(login["name"])
    signup_email = driver.find_element(By.CSS_SELECTOR, "input[data-qa='signup-email']")
    signup_email.clear()
    signup_email.send_keys(login["email"])
    safe_click(driver, driver.find_element(By.CSS_SELECTOR, "button[data-qa='signup-button']"))

    wait.until(EC.presence_of_element_located((By.ID, "password"))).send_keys(login["password"])

    driver.find_element(By.ID, "days").send_keys("1")
    driver.find_element(By.ID, "months").send_keys("January")
    driver.find_element(By.ID, "years").send_keys("2000")
    driver.find_element(By.ID, "first_name").send_keys(login["name"].split(" ")[0])
    driver.find_element(By.ID, "last_name").send_keys(
        login["name"].split(" ")[-1] if " " in login["name"] else "User"
    )
    driver.find_element(By.ID, "address1").send_keys("221B Test Street")
    driver.find_element(By.ID, "state").send_keys("West Bengal")
    driver.find_element(By.ID, "city").send_keys("Kolkata")
    driver.find_element(By.ID, "zipcode").send_keys("700001")
    driver.find_element(By.ID, "mobile_number").send_keys("9876543210")

    safe_click(driver, driver.find_element(By.CSS_SELECTOR, "button[data-qa='create-account']"))

    wait.until(EC.presence_of_element_located((By.CSS_SELECTOR, "a[data-qa='continue-button']")))
    safe_click(driver, driver.find_element(By.CSS_SELECTOR, "a[data-qa='continue-button']"))
    wait.until(EC.presence_of_element_located((By.XPATH, "//a[contains(., 'Logged in as')]")))


# Search product
def search_product(driver, base_url, search_term):
    """Searches on the Products page and returns the list of result cards."""
    wait = WebDriverWait(driver, WAIT_TIMEOUT)
    driver.get(f"{base_url}/products")

    search_box = wait.until(EC.presence_of_element_located((By.ID, "search_product")))
    search_box.clear()
    search_box.send_keys(search_term)
    safe_click(driver, driver.find_element(By.ID, "submit_search"))

    wait.until(EC.presence_of_element_located((By.CLASS_NAME, "features_items")))
    return driver.find_elements(By.CSS_SELECTOR, ".features_items .product-image-wrapper")


# Add product to cart + update quantity
def add_product_with_quantity(driver, product_name, quantity):
    wait = WebDriverWait(driver, WAIT_TIMEOUT)

    product_link = driver.find_element(
        By.XPATH,
        f"//p[text()='{product_name}']/ancestor::div[@class='product-image-wrapper']"
        f"//a[contains(text(), 'View Product')]",
    )
    driver.get(product_link.get_attribute("href"))

    wait.until(EC.presence_of_element_located((By.ID, "quantity")))
    qty_field = driver.find_element(By.ID, "quantity")
    qty_field.clear()
    qty_field.send_keys(str(quantity))
    entered_quantity = qty_field.get_attribute("value")

    safe_click(driver, driver.find_element(By.CSS_SELECTOR, "button.cart"))

    dismiss_browser_alert_if_present(driver)
    close_add_to_cart_modal(driver)
    return entered_quantity


# Read cart details
def get_cart_quantity(driver, base_url, product_name):
    wait = WebDriverWait(driver, WAIT_TIMEOUT)
    driver.get(f"{base_url}/view_cart")
    wait.until(lambda d: d.find_elements(By.ID, "cart_info_table")
               or d.find_elements(By.ID, "empty_cart"))
    if not driver.find_elements(By.ID, "cart_info_table"):
        return None  # cart is empty

    rows = driver.find_elements(
        By.XPATH, f"//tbody/tr[td[@class='cart_description']//a[text()='{product_name}']]"
    )
    if not rows:
        return None
    return rows[0].find_element(By.CSS_SELECTOR, ".cart_quantity button").text.strip()


def clear_cart(driver, base_url):
    wait = WebDriverWait(driver, WAIT_TIMEOUT)
    driver.get(f"{base_url}/view_cart")
    wait.until(lambda d: d.find_elements(By.ID, "cart_info_table")
               or d.find_elements(By.ID, "empty_cart"))

    while True:
        delete_buttons = driver.find_elements(By.CSS_SELECTOR, "a.cart_quantity_delete")
        if not delete_buttons:
            break
        first_button = delete_buttons[0]
        safe_click(driver, first_button)
        wait.until(EC.staleness_of(first_button))